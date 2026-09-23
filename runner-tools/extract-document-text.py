#!/usr/bin/env python3
"""Extract bounded, readable text from common untrusted document attachments."""

from __future__ import annotations

import argparse
import csv
import html.parser
import io
import pathlib
import subprocess
import sys
import xml.etree.ElementTree as element_tree
import zipfile

MAX_OUTPUT_CHARACTERS = 2_000_000
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024


class _HtmlTextExtractor(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def _extract_pdf(path: pathlib.Path) -> str:
    completed = subprocess.run(
        ["pdftotext", "-layout", "-nopgbrk", str(path), "-"],
        check=False,
        capture_output=True,
        timeout=120,
    )
    if completed.returncode == 0:
        return completed.stdout.decode("utf-8", errors="replace")

    from pypdf import PdfReader

    return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _extract_docx(path: pathlib.Path) -> str:
    _validate_office_archive(path)
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: pathlib.Path) -> str:
    _validate_office_archive(path)
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    output = io.StringIO()
    for sheet in workbook.worksheets:
        output.write(f"\n## Sheet: {sheet.title}\n")
        writer = csv.writer(output, delimiter="\t", lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            writer.writerow("" if value is None else str(value) for value in row)
    workbook.close()
    return output.getvalue()


def _extract_pptx(path: pathlib.Path) -> str:
    _validate_office_archive(path)
    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(
            name for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        for number, slide in enumerate(slides, start=1):
            root = element_tree.fromstring(archive.read(slide))
            text = " ".join(
                node.text or "" for node in root.findall(".//a:t", namespace)
            ).strip()
            parts.append(f"## Slide {number}\n{text}")
    return "\n\n".join(parts)


def _validate_office_archive(path: pathlib.Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_ENTRIES:
            raise ValueError("Office archive contains too many entries")
        total_size = sum(member.file_size for member in members)
        if total_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
            raise ValueError("Office archive expands beyond the 200 MiB safety limit")
        if any(member.flag_bits & 0x1 for member in members):
            raise ValueError("Encrypted Office archives are not supported")


def _extract_html(path: pathlib.Path) -> str:
    parser = _HtmlTextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return "\n".join(parser.parts)


def _extract_text(path: pathlib.Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".docm"}:
        return _extract_docx(path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _extract_xlsx(path)
    if suffix in {".pptx", ".pptm", ".ppsx", ".ppsm"}:
        return _extract_pptx(path)
    if suffix in {".html", ".htm"}:
        return _extract_html(path)
    if suffix in {
        ".txt", ".md", ".rst", ".rtf", ".log", ".csv", ".tsv", ".json",
        ".jsonl", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
        ".eml",
    }:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(
        f"Unsupported document type '{suffix or '(none)'}'. "
        "Supported types: PDF, DOCX, XLSX, PPTX, HTML, and text-based documents."
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract text from a PDF, Office, HTML, or text document without executing it."
    )
    parser.add_argument("path", type=pathlib.Path)
    arguments = parser.parse_args()
    path = arguments.path.resolve(strict=True)
    if not path.is_file():
        parser.error("path must refer to a regular file")

    try:
        text = _extract_text(path)
    except (OSError, ValueError, zipfile.BadZipFile) as exception:
        print(f"extract-document-text: {exception}", file=sys.stderr)
        return 2

    if len(text) > MAX_OUTPUT_CHARACTERS:
        text = text[:MAX_OUTPUT_CHARACTERS] + "\n\n[Output truncated at 2,000,000 characters.]\n"
    sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
